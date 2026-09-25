"""Parse pasted GPT account rows without logging their credentials."""
from __future__ import annotations

import csv
from io import BytesIO
import re
from urllib.parse import parse_qs, urlparse
from zipfile import BadZipFile, ZipFile


_EMAIL = re.compile(r"^[^\s@|,]+@[^\s@|,]+\.[^\s@|,]+$")
_SEPARATOR = re.compile(r"^:?-+:?$")


def valid_email(value: str) -> bool:
    return bool(_EMAIL.fullmatch(str(value or "").strip()))


def _registered_account_from_parts(parts: list[str]) -> dict | None:
    if len(parts) < 4 or not valid_email(parts[0]):
        return None
    secret = re.sub(r"^\s*2fa\s*[:：]\s*", "", parts[2], flags=re.IGNORECASE).strip()
    if secret.startswith("otpauth://"):
        secret = (parse_qs(urlparse(secret).query).get("secret") or [""])[0].strip()
    if re.fullmatch(r"[A-Z2-7= ]+", secret, flags=re.IGNORECASE):
        secret = secret.replace(" ", "")
    if secret.isdigit() and len(secret) == 6:
        return None
    return {
        "email": parts[0],
        "registration_password": parts[1],
        "totp_secret": secret,
        "access_token": parts[3],
    }


def parse_registered_accounts(text: str) -> tuple[list[dict], int]:
    """Accept Markdown tables, TSV/CSV, pipes and ``----`` separated rows."""
    records: list[dict] = []
    invalid = 0
    for raw in str(text or "").splitlines():
        line = raw.strip().lstrip("\ufeff")
        if not line or line.startswith("#"):
            continue
        if "|" in line:
            parts = [part.strip() for part in line.strip("|").split("|")]
        elif "\t" in line:
            parts = [part.strip() for part in line.split("\t")]
        elif "----" in line:
            parts = [part.strip() for part in line.split("----", 3)]
        else:
            parts = [part.strip() for part in next(csv.reader([line]))]
        if len(parts) >= 4 and (
            parts[0].lower() in {"邮箱", "email", "e-mail"}
            or all(_SEPARATOR.fullmatch(part) for part in parts[:4])
        ):
            continue
        record = _registered_account_from_parts(parts)
        if record is None:
            invalid += 1
            continue
        records.append(record)
    return records, invalid


_XLSX_HEADERS = {
    "email": {"邮箱", "邮箱地址", "email", "emailaddress"},
    "password": {"密码", "gpt密码", "chatgpt密码", "password"},
    "totp": {"2fa", "2fa密钥", "totp", "totpsecret"},
    "token": {"at", "accesstoken", "token"},
}


def _header_columns(values: tuple) -> dict[str, int] | None:
    columns = {}
    for index, value in enumerate(values):
        name = re.sub(r"[\s_\-/]+", "", str(value or "").strip()).lower()
        for field, names in _XLSX_HEADERS.items():
            if name in names and field not in columns:
                columns[field] = index
    return columns if len(columns) == 4 else None


def parse_registered_accounts_xlsx(data: bytes) -> tuple[list[dict], int]:
    """Read account columns from a bounded .xlsx archive; never evaluate formulas."""
    if not data or len(data) > 8_000_000:
        raise ValueError("Excel 文件不能为空或超过 8 MB")
    try:
        with ZipFile(BytesIO(data)) as archive:
            members = archive.infolist()
            if len(members) > 500 or sum(item.file_size for item in members) > 80_000_000:
                raise ValueError("Excel 文件内容过大")
    except BadZipFile as exc:
        raise ValueError("文件不是有效的 .xlsx") from exc

    from openpyxl import load_workbook

    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
    except Exception as exc:
        raise ValueError("无法读取 Excel 文件") from exc
    records: list[dict] = []
    invalid = 0
    matched_sheet = False
    try:
        for sheet in workbook.worksheets:
            if sheet.max_row and sheet.max_row > 50_000:
                raise ValueError("Excel 工作表超过 50000 行")
            width = min(sheet.max_column or 0, 30)
            if width < 4:
                continue
            header = None
            for row_number, row in enumerate(sheet.iter_rows(max_row=20, max_col=width, values_only=True), 1):
                columns = _header_columns(row)
                if columns:
                    header = (row_number, columns)
                    break
            if header is None:
                continue
            matched_sheet = True
            start_row, columns = header
            for row in sheet.iter_rows(min_row=start_row + 1, max_col=width):
                cells = [row[columns[field]] for field in ("email", "password", "totp", "token")]
                if all(cell.value is None or str(cell.value).strip() == "" for cell in cells):
                    continue
                if any(cell.data_type == "f" or
                       (cell.value is not None and not isinstance(cell.value, str)) for cell in cells):
                    invalid += 1
                    continue
                parts = [str(cell.value or "").strip() for cell in cells]
                record = _registered_account_from_parts(parts)
                if record is None:
                    invalid += 1
                else:
                    records.append(record)
    finally:
        workbook.close()
    if not matched_sheet:
        raise ValueError("找不到“邮箱、密码、2FA、AT”四列表头")
    return records, invalid
